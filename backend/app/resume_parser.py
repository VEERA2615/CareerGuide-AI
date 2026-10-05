from __future__ import annotations

import io
import re

from docx import Document
from pypdf import PdfReader


SKILL_ALIASES = {
    "Python": ["python"],
    "SQL": ["sql", "mysql", "postgresql", "postgres"],
    "Excel": ["excel", "microsoft excel"],
    "Power BI": ["power bi"],
    "Tableau": ["tableau"],
    "Pandas": ["pandas"],
    "NumPy": ["numpy"],
    "Matplotlib": ["matplotlib"],
    "Seaborn": ["seaborn"],
    "Statistics": ["statistics", "statistical"],
    "Machine Learning": ["machine learning", "machine-learning", "ml"],
    "Scikit-learn": ["scikit-learn", "sklearn"],
    "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"],
    "Java": ["java"],
    "JavaScript": ["javascript", "js"],
    "React": ["react", "react.js"],
    "Node.js": ["node.js", "nodejs", "node"],
    "HTML/CSS": ["html", "css"],
    "Git/GitHub": ["git", "github"],
    "REST APIs": ["rest api", "rest apis", "api"],
    "ETL": ["etl", "extract transform load"],
    "Data Visualization": ["data visualization", "data visualisation"],
    "Data Cleaning": ["data cleaning", "data preprocessing"],
}


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_resume_text(filename: str, content_type: str, data: bytes) -> str:
    lower = filename.lower()

    if lower.endswith(".pdf") or "pdf" in content_type:
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            return clean_text("\n\n".join(pages))
        except Exception as exc:
            raise ValueError(f"Could not read the PDF: {type(exc).__name__}") from exc

    if lower.endswith(".docx") or "wordprocessingml" in content_type:
        try:
            document = Document(io.BytesIO(data))
            paragraphs = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            return clean_text("\n".join(paragraphs))
        except Exception as exc:
            raise ValueError(f"Could not read the DOCX: {type(exc).__name__}") from exc

    if lower.endswith(".txt") or content_type.startswith("text/"):
        try:
            return clean_text(data.decode("utf-8", errors="ignore"))
        except Exception as exc:
            raise ValueError("Could not read the text file.") from exc

    raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT resume.")


def detect_skills(text: str) -> list[str]:
    haystack = text.lower()
    found = []
    for display, aliases in SKILL_ALIASES.items():
        if any(re.search(r"(?<!\w)" + re.escape(alias) + r"(?!\w)", haystack) for alias in aliases):
            found.append(display)
    return sorted(found)
